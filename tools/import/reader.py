"""Step 1 (T-004): reads `Example-Master.xlsx` without ever writing back to it.

Reads column by fixed letter (not by matching the header text at runtime) because the sheet has a
stray unlabelled column (Task Management column N) and a documented "~" filler column (Ver. 1.1
column AB) that would otherwise collide when zipping header cells with data cells. `expected_header`
on each column is checked against row 4 so a change to the workbook's layout fails loudly instead of
silently reading the wrong column.

Only `docs/05-import-spec.md`'s two "in scope" sheets are read row by row: `Ver. 1.1` and
`Task Management`. `Ver. 1.0`, `BACKUP`, `Feeding Management` are listed in the sheet inventory
(for the report) but never opened row by row (OQ-7).
"""

from __future__ import annotations

import hashlib
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.worksheet.worksheet import Worksheet

VER11_SHEET = "Ver. 1.1"
TASK_MANAGEMENT_SHEET = "Task Management"
NOT_IMPORTED_SHEETS = ("Ver. 1.0", "BACKUP", "Feeding Management")

# (column letter, exact header text in row 4, key in the raw record). Column AB ("~") is a filler
# column with no mapping (docs/05-import-spec.md section 2) and is intentionally left out.
VER11_COLUMNS: tuple[tuple[str, str, str], ...] = (
    ("B", "Check", "check"),
    ("C", "No.", "no"),
    ("D", "LINE", "line"),
    ("E", "Add Info 01", "gene_raw"),
    ("F", "Resouces01", "source"),
    ("G", "DOB (Current)", "dob"),
    ("H", "Status", "status"),
    ("I", "ID Method", "id_method"),
    ("J", "ID Date", "id_date"),
    ("K", "ID Number", "id_number"),
    ("L", "Primer F", "primer_f_name"),
    ("M", "Sequence F", "primer_f_seq"),
    ("N", "Primer R", "primer_r_name"),
    ("O", "Sequence R", "primer_r_seq"),
    ("P", "Base Pairs / b. p.", "base_pairs"),
    ("Q", "PCR Cond. / C", "pcr_cond"),
    ("R", "Seq. Cond.", "seq_cond"),
    ("S", "Mutation", "mutation"),
    ("T", "Guide Sequence", "guide_seq"),
    ("U", "Phenotype", "phenotype"),
    ("V", "Resource02", "resource02"),
    ("W", "Cryo Preservation", "cryo_flag"),
    ("X", "Cryo Date", "cryo_date"),
    ("Y", "Cryo Place", "cryo_place"),
    ("Z", "Cryo BoxID", "cryo_box"),
    ("AA", "CryoID_min", "cryo_id_min"),
    ("AC", "CryoID_MAX", "cryo_id_max"),
    ("AD", "Add Info 02", "add_info_02"),
    ("AE", "Last Update", "last_update"),
)

TASK_MANAGEMENT_COLUMNS: tuple[tuple[str, str, str], ...] = (
    ("C", "Posted", "posted"),
    ("D", "Category", "category"),
    ("E", "Task", "task"),
    ("F", "Line", "line"),
    ("G", "Description", "description"),
    ("H", "Assigned to", "assigned_to"),
    ("I", "Due", "due"),
    ("J", "Status", "status"),
    ("K", "Message", "message"),
    ("L", "Last Updated", "last_updated"),
)

HEADER_ROW = 4
FIRST_DATA_ROW = 5


@dataclass(frozen=True)
class SheetInfo:
    name: str
    state: str  # "visible" | "hidden" | "veryHidden"
    max_row: int


@dataclass(frozen=True)
class RawRecord:
    row: int
    row_hidden: bool
    values: dict[str, Any]


def file_fingerprint(path: Path) -> tuple[float, str]:
    """`(mtime, sha256)`. Tests call this before and after a read to prove the file is untouched."""
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    return path.stat().st_mtime, digest


def _check_headers(sheet: Worksheet, columns: tuple[tuple[str, str, str], ...]) -> None:
    for letter, expected, _key in columns:
        actual = sheet[f"{letter}{HEADER_ROW}"].value
        if (actual or "").strip() != expected:
            raise ValueError(
                f"{sheet.title}!{letter}{HEADER_ROW}: expected header {expected!r}, found {actual!r} "
                "-- the workbook layout may have changed; update reader.py's column list."
            )


def _read_rows(
    sheet: Worksheet,
    columns: tuple[tuple[str, str, str], ...],
    first_row: int,
    last_row: int,
    line_column: str,
) -> list[RawRecord]:
    _check_headers(sheet, columns)
    records: list[RawRecord] = []
    for row in range(first_row, last_row + 1):
        line_value = sheet[f"{line_column}{row}"].value
        if line_value is None or (isinstance(line_value, str) and line_value.strip() == ""):
            continue
        values = {key: sheet[f"{letter}{row}"].value for letter, _header, key in columns}
        records.append(
            RawRecord(row=row, row_hidden=bool(sheet.row_dimensions[row].hidden), values=values)
        )
    return records


def read_sheet_inventory(path: Path) -> list[SheetInfo]:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        return [
            SheetInfo(name=name, state=workbook[name].sheet_state, max_row=workbook[name].max_row)
            for name in workbook.sheetnames
        ]
    finally:
        workbook.close()


def read_ver11_rows(path: Path) -> list[RawRecord]:
    """Every `Ver. 1.1` row with a non-empty `LINE` cell, rows 5 onward, hidden or not."""
    workbook = openpyxl.load_workbook(path, data_only=True)
    try:
        sheet = workbook[VER11_SHEET]
        return _read_rows(sheet, VER11_COLUMNS, FIRST_DATA_ROW, sheet.max_row, line_column="D")
    finally:
        workbook.close()


def read_task_management_rows(path: Path) -> list[RawRecord]:
    """Every `Task Management` row with a non-empty `Line` cell, rows 5 onward, hidden or not."""
    workbook = openpyxl.load_workbook(path, data_only=True)
    try:
        sheet = workbook[TASK_MANAGEMENT_SHEET]
        return _read_rows(sheet, TASK_MANAGEMENT_COLUMNS, FIRST_DATA_ROW, sheet.max_row, line_column="F")
    finally:
        workbook.close()
