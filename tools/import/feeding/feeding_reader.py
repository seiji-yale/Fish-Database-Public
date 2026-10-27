"""Reads the lab's feeding Google Sheet, downloaded as `.xlsx` (T-036, docs/08 section 7.1).

One tab per month, named `MonthYYYY` (`October2026`). In each tab the weekday headers are in `C5:I5`, the
year is in `D4`, and the month is laid out in week blocks of three rows: the day numbers (formulas, so the
file must hold their cached values), the feeder's first name, and an event text. The reader takes the month
from the tab name and checks the layout against it instead of trusting the sheet's own month selector (which
is missing or wrong in some tabs). It never writes to the file, and it does not depend on a tab being hidden
or visible. A layout it does not recognise raises `LayoutError`: it never guesses (T-004 lesson).
"""

from __future__ import annotations

import calendar
import re
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path

import openpyxl

MONTH_NAMES = [calendar.month_name[number] for number in range(1, 13)]
TAB_NAME = re.compile(r"^(" + "|".join(MONTH_NAMES) + r")(\d{4})$")
WEEKDAY_HEADERS = ["Sun", "Mon", "Tue", "Wed", "Thu", "Fri", "Sat"]
FIRST_COLUMN = 3  # C
FIRST_ROW, LAST_ROW = 6, 40


class LayoutError(Exception):
    """The workbook does not look like the feeding sheet; the message says where and what to do."""


@dataclass(frozen=True)
class Assignment:
    date: str  # YYYY-MM-DD
    name: str  # the feeder's name as written in the sheet (trimmed)
    cell: str  # e.g. `October2026!E7`, kept in `imported_from`


@dataclass(frozen=True)
class Orphan:
    cell: str
    name: str  # a name with no day number above it


@dataclass
class MonthRead:
    month: str  # YYYY-MM
    tab: str
    days: int = 0  # day cells of this month in the grid
    assignments: list[Assignment] = field(default_factory=list)
    orphans: list[Orphan] = field(default_factory=list)
    events_ignored: int = 0


def tab_month(tab: str) -> str | None:
    """`October2026` -> `2026-10`; None when the tab is not a month tab."""
    match = TAB_NAME.match(tab)
    if match is None:
        return None
    return f"{match.group(2)}-{MONTH_NAMES.index(match.group(1)) + 1:02d}"


def _text(value: object) -> str:
    return " ".join(str(value).split()) if value is not None else ""


def _day_number(value: object) -> int | None:
    """A cached day number (formulas are saved as numbers); text that is not one is not a day."""
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)) and float(value).is_integer() and 1 <= int(value) <= 31:
        return int(value)
    return None


def read_month(sheet, tab: str, month: str) -> MonthRead:
    year, number = (int(part) for part in month.split("-"))
    headers = [_text(sheet.cell(5, FIRST_COLUMN + offset).value) for offset in range(7)]
    if headers != WEEKDAY_HEADERS:
        raise LayoutError(
            f"{tab}: C5:I5 should be {', '.join(WEEKDAY_HEADERS)} but is {', '.join(headers)}. "
            "The sheet layout changed; the reader was not told about it."
        )
    sheet_year = sheet["D4"].value
    if not isinstance(sheet_year, (int, float)) or int(sheet_year) != year:
        raise LayoutError(f"{tab}: D4 should hold the year {year} but holds {sheet_year!r}.")
    first_weekday = (date(year, number, 1).weekday() + 1) % 7  # Sunday = 0
    days_in_month = calendar.monthrange(year, number)[1]

    read = MonthRead(month=month, tab=tab)
    day_rows = [
        row
        for row in range(FIRST_ROW, LAST_ROW + 1)
        if any(_day_number(sheet.cell(row, FIRST_COLUMN + offset).value) is not None for offset in range(7))
    ]
    if not day_rows:
        raise LayoutError(
            f"{tab}: no day numbers found in rows {FIRST_ROW}-{LAST_ROW}. The file may have been saved "
            "without the formulas' values: download it again from Google Sheets as .xlsx."
        )
    seen: list[int] = []
    for row in day_rows:
        for offset in range(7):
            cell = sheet.cell(row, FIRST_COLUMN + offset)
            day = _day_number(cell.value)
            name = _text(sheet.cell(row + 1, FIRST_COLUMN + offset).value)
            name_cell = f"{tab}!{sheet.cell(row + 1, FIRST_COLUMN + offset).coordinate}"
            if _text(sheet.cell(row + 2, FIRST_COLUMN + offset).value) != "":
                read.events_ignored += 1
            if day is None:
                if name != "":
                    read.orphans.append(Orphan(cell=name_cell, name=name))
                continue
            seen.append(day)
            read.days += 1
            if name != "":
                read.assignments.append(Assignment(f"{year:04d}-{number:02d}-{day:02d}", name, name_cell))
    if seen != list(range(1, days_in_month + 1)):
        raise LayoutError(
            f"{tab}: the day numbers are not 1 to {days_in_month} in order ({seen[:8]}...). "
            "The sheet's month does not match its tab."
        )
    first_row = day_rows[0]
    first_column = next(
        offset for offset in range(7) if _day_number(sheet.cell(first_row, FIRST_COLUMN + offset).value) == 1
    )
    if first_column != first_weekday:
        raise LayoutError(
            f"{tab}: day 1 is in the {WEEKDAY_HEADERS[first_column]} column but {month}-01 is a "
            f"{WEEKDAY_HEADERS[first_weekday]}."
        )
    return read


def read_workbook(path: Path, months: list[str] | None = None) -> list[MonthRead]:
    """The requested months (all month tabs when `months` is None), in date order. A requested month with
    no tab raises `LayoutError`."""
    workbook = openpyxl.load_workbook(path, data_only=True)
    tabs = {tab_month(sheet.title): sheet for sheet in workbook.worksheets if tab_month(sheet.title)}
    wanted = sorted(tabs) if months is None else sorted(months)
    missing = [month for month in wanted if month not in tabs]
    if missing:
        raise LayoutError(f"No tab for {', '.join(missing)}; the workbook has {', '.join(sorted(tabs))}.")
    return [read_month(tabs[month], tabs[month].title, month) for month in wanted]
