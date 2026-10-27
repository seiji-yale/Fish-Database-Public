"""A synthetic feeding workbook for the tests (never the real sheet): month tabs in the layout of docs/08
section 7.1, with the day numbers saved as plain numbers (what a Google Sheets export holds as the cached
values of the sheet's formulas)."""

from __future__ import annotations

import calendar
from datetime import date
from pathlib import Path

import openpyxl
import pytest

MONTH_ROWS = [6, 11, 16, 21, 26, 32]
COLUMNS = "CDEFGHI"


def add_month_tab(
    workbook,
    year: int,
    month: int,
    names: dict[int, str],
    *,
    events: dict[int, str] | None = None,
    orphans: dict[str, str] | None = None,
    title: str | None = None,
    headers: list[str] | None = None,
    year_cell: object = None,
    skip_days: bool = False,
) -> None:
    """`names` maps a day of the month to the feeder's name; `orphans` maps a cell (`C7`) to a name that
    has no day number above it."""
    sheet = workbook.create_sheet(title or f"{calendar.month_name[month]}{year}")
    sheet["D4"] = year if year_cell is None else year_cell
    for offset, text in enumerate(headers or ["Sun", "Mon", "Tue", "Wed", "Thu", "Fri", "Sat"]):
        sheet.cell(5, 3 + offset, text)
    first = (date(year, month, 1).weekday() + 1) % 7
    total = calendar.monthrange(year, month)[1]
    slot = first
    for day in range(1, total + 1):
        row = MONTH_ROWS[slot // 7]
        column = 3 + slot % 7
        if not skip_days:
            sheet.cell(row, column, day)
        if day in names:
            sheet.cell(row + 1, column, names[day])
        if events and day in events:
            sheet.cell(row + 2, column, events[day])
        slot += 1
    for cell, name in (orphans or {}).items():
        sheet[cell] = name
    sheet["C41"] = "Emergency contact sentence (not imported)"


@pytest.fixture
def workbook_path(tmp_path: Path) -> Path:
    """October 2026 (an orphan before the 1st, an event, one unknown name) and November 2026 (a partial
    month: only the first week is filled, like the template leftovers)."""
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    october = {day: ("Feeder One" if day % 2 else "Feeder Two") for day in range(1, 32)}
    october[3] = "Newcomer"  # a name no user has
    october[10] = "  feeder two  "  # case and spaces do not matter
    add_month_tab(
        workbook, 2026, 10, october, events={5: "Test Water Quality"}, orphans={"E7": "Feeder One", "F7": "Feeder One"}
    )
    add_month_tab(workbook, 2026, 11, {2: "Feeder Three", 3: "Feeder Three", 4: "Feeder Three", 5: "Feeder Three", 6: "Feeder Three", 7: "Feeder Three"})
    add_month_tab(workbook, 2026, 9, {})  # an empty month
    path = tmp_path / "feeding.xlsx"
    workbook.save(path)
    return path
